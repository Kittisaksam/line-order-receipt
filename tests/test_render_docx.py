from datetime import date
from decimal import Decimal

from docx import Document

from arma.billing import Bill, BillLine
from arma.render_docx import DocxRenderer

SETTINGS = {
    "shop": {"name": "ARMA LASER", "address_lines": ["ที่อยู่ 1"]},
    "customer": {"name": "บริษัท ลูกค้า จำกัด", "address_lines": ["ที่ตั้ง 2"], "tax_id": "0105500000000"},
    "billing": {"terms": "ครบกำหนดเมื่อได้รับ"},
}


def _bills() -> list[Bill]:
    a = Bill(no="337", lines=[
        BillLine(date(2026, 10, 2), 3, "LA-A", "ตะแกรง FJ สีแดง", Decimal(700)),
        BillLine(date(2026, 10, 3), 10, "LA-B", "TOYOTA ท้ายรถ สีขาว", Decimal(170)),
    ])
    b = Bill(no="338", lines=[BillLine(date(2026, 10, 13), 1, "SV-DESIGN", "งานออกแบบ EASY PASS", Decimal(200))])
    return [a, b]


def _text(doc) -> str:
    cells = [c.text for t in doc.tables for r in t.rows for c in r.cells]
    return "\n".join([p.text for p in doc.paragraphs] + cells)


def _items_table(doc):
    return next(t for t in doc.tables if "คำอธิบาย" in [c.text for c in t.rows[0].cells])


def test_receipt_docx(tmp_path):
    out = DocxRenderer(SETTINGS).render(_bills()[:1], tmp_path / "r.docx")
    doc = Document(out)
    text = _text(doc)
    assert "ใบเสร็จรับเงิน #337" in text and "เลขประจำตัวผู้เสียภาษี : 0105500000000" in text
    items = _items_table(doc)
    assert len(items.rows) == 3  # หัว + 2 รายการ
    assert [c.text for c in items.rows[1].cells] == ["3", "ตะแกรง FJ สีแดง", "700", "2,100"]
    assert "3,800" in text  # ยอดรวม


def test_merged_receipts_one_page_per_bill(tmp_path):
    doc = Document(DocxRenderer(SETTINGS).render(_bills(), tmp_path / "all.docx"))
    text = _text(doc)
    assert "ใบเสร็จรับเงิน #337" in text and "ใบเสร็จรับเงิน #338" in text
    assert doc.element.body.xml.count('w:type="page"') == 1


def test_summary_docx(tmp_path):
    doc = Document(DocxRenderer(SETTINGS).render_summary(_bills(), 2026, 10, tmp_path / "s.docx"))
    text = _text(doc)
    assert "ใบสรุปรายการประจำเดือน" in text and "ครอบคลุมใบเสร็จ #337 – #338" in text
    items = _items_table(doc)
    assert items.rows[0].cells[0].text == "ใบเสร็จ" and len(items.rows) == 4
    assert [c.text for c in items.rows[3].cells][:2] == ["#338", "-"]  # ค่าบริการแสดงจำนวนเป็น "-"
    assert "ยอดรวมทั้งเดือน" in text and "4,000" in text
