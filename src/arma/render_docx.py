"""ขั้นที่ 7 (Word) — ใบเสร็จ/ใบสรุปเดือนเป็น .docx ให้ admin แก้เองได้ (หน้าตาตาม receipt.html.j2)"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from .billing import Bill
from .thai_utils import bahttext, fmt_date_th, fmt_money, month_name_th

LEFT, CENTER, RIGHT = WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT
BASE_PT = 10
GREY = RGBColor(0x55, 0x55, 0x55)


def _style_run(run, size: float = BASE_PT, bold: bool = False, color: RGBColor | None = None) -> None:
    # ภาษาไทยเป็น complex script → ต้องตั้ง szCs/bCs ด้วย ไม่งั้น Word ใช้ขนาด/ตัวหนาเริ่มต้น
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    # วางต่อจาก w:sz / w:b ตามลำดับ schema (Word ฟ้องไฟล์เสียถ้าลำดับผิด)
    rpr = run._r.get_or_add_rPr()
    sz = OxmlElement("w:szCs")
    sz.set(qn("w:val"), str(int(size * 2)))
    rpr.find(qn("w:sz")).addnext(sz)
    if bold:
        rpr.find(qn("w:b")).addnext(OxmlElement("w:bCs"))


def _para(container, text: str = "", align=LEFT, size: float = BASE_PT, bold: bool = False,
          color: RGBColor | None = None, space_after: float = 0):
    p = container.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    if text:
        _style_run(p.add_run(text), size, bold, color)
    return p


def _cell(cell, text: str, align=LEFT, bold: bool = False, size: float = BASE_PT) -> None:
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    if text:
        _style_run(p.add_run(text), size, bold)


def _widths(table, widths_mm: list[float]) -> None:
    # ใส่ tblW/gridCol ตรง ๆ — Pages/Google Docs/Quick Look ไม่ได้อ่านแค่ความกว้างของ cell
    table.autofit = False
    tblpr = table._tbl.tblPr
    for old in tblpr.findall(qn("w:tblW")):
        tblpr.remove(old)
    tblw = OxmlElement("w:tblW")
    tblw.set(qn("w:w"), str(int(sum(widths_mm) * 56.7)))
    tblw.set(qn("w:type"), "dxa")
    style = tblpr.find(qn("w:tblStyle"))
    (style.addnext if style is not None else lambda e: tblpr.insert(0, e))(tblw)
    for col, w in zip(table.columns, widths_mm):
        col.width = Mm(w)
    for row in table.rows:
        for c, w in zip(row.cells, widths_mm):
            c.width = Mm(w)


def _grid(table) -> None:
    """เส้นตารางทุกช่อง (ตั้งตรง ๆ ไม่พึ่ง style "Table Grid" อย่างเดียว)"""
    table.style = "Table Grid"
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "444444")
        borders.append(el)
    tblpr = table._tbl.tblPr
    # ลำดับ schema: tblStyle, …, tblW, jc, …, tblBorders, …, tblLayout
    # (element ของ lxml ที่ไม่มีลูกเป็น False → เช็ก None ตรง ๆ)
    anchor = next(e for e in (tblpr.find(qn(t)) for t in ("w:jc", "w:tblW", "w:tblStyle")) if e is not None)
    anchor.addnext(borders)


def _repeat_header(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trpr.append(el)


class DocxRenderer:
    def __init__(self, settings: dict):
        self.settings = settings
        self.font = settings.get("pdf", {}).get("font_family", "Sarabun")

    # ---------- เอกสาร ----------
    def _new_doc(self) -> Document:
        doc = Document()
        sec = doc.sections[0]
        sec.page_width, sec.page_height = Mm(210), Mm(297)
        sec.top_margin, sec.bottom_margin, sec.left_margin, sec.right_margin = Mm(16), Mm(14), Mm(14), Mm(14)
        normal = doc.styles["Normal"]
        normal.font.name = self.font
        normal.font.size = Pt(BASE_PT)
        fonts = normal.element.get_or_add_rPr().get_or_add_rFonts()
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            fonts.set(qn(attr), self.font)
        normal.paragraph_format.space_after = Pt(0)
        return doc

    def render(self, bills: list[Bill], out: Path) -> Path:
        """ใบเสร็จ 1 ไฟล์ (หลายบิลได้ คั่นด้วยขึ้นหน้าใหม่)"""
        doc = self._new_doc()
        for i, bill in enumerate(bills):
            if i:
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            self._receipt(doc, bill)
        return self._save(doc, out)

    def render_summary(self, bills: list[Bill], year: int, month: int, out: Path) -> Path:
        """ใบสรุปทั้งเดือน (ไม่ใช่ใบเสร็จ) — ตรงกับ Renderer.render_summary"""
        doc = self._new_doc()
        everything = Bill(lines=[ln for b in bills for ln in b.lines], time=bills[0].time, date_rule="last")
        numbers = [b.no for b in bills]
        covers = f"#{numbers[0]}" if len(numbers) == 1 else f"#{numbers[0]} – #{numbers[-1]}"
        self._header(doc, "ใบสรุปรายการประจำเดือน", [
            f"เดือน{month_name_th(year, month)}",
            f"ครอบคลุมใบเสร็จ {covers}",
            f"ช่วงวันที่ {fmt_date_th(min(b.date for b in bills))} – {fmt_date_th(everything.date)}",
        ])
        self._parties(doc)
        p = _para(doc, "ข้อคิดเห็นหรือคำแนะนำพิเศษ: ", bold=True, space_after=10)
        p.paragraph_format.space_before = Pt(18)
        _style_run(p.add_run("เอกสารสรุปเพื่อตรวจสอบ ไม่ใช่ใบเสร็จรับเงิน"), color=RGBColor(0xAA, 0x33, 0x33))
        self._items(doc, [(b.no, ln) for b in bills for ln in b.lines], with_billno=True)
        rows = [(f"ใบเสร็จ #{b.no} ({fmt_date_th(b.date)})", fmt_money(b.subtotal), False) for b in bills]
        rows.append(("ยอดรวมทั้งเดือน", fmt_money(everything.subtotal), True))
        self._totals(doc, rows)
        return self._save(doc, out)

    # ---------- ส่วนประกอบ ----------
    def _receipt(self, doc, bill: Bill) -> None:
        billing = self.settings.get("billing", {})
        date_th = fmt_date_th(bill.date)
        self._header(doc, "ใบเสร็จรับเงิน", [f"ใบเสร็จรับเงิน #{bill.no}", f"เวลา:{bill.time}น. วันที่ {date_th}"])
        self._parties(doc)
        p = _para(doc, "ข้อคิดเห็นหรือคำแนะนำพิเศษ:", bold=True, space_after=24)
        p.paragraph_format.space_before = Pt(18)

        terms = doc.add_table(rows=2, cols=6)
        for c, h in zip(terms.rows[0].cells, ["พนักงานขาย", "P.O. หมายเลข", "ผู้จัดหาวัตถุดิบ",
                                              "ส่งผ่านทาง", "F.O.B. จุด", "ข้อตกลง"]):
            _cell(c, h, CENTER, bold=True)
        _cell(terms.rows[1].cells[5], billing.get("terms", "ครบกำหนดเมื่อได้รับ"), size=BASE_PT - 1)
        terms.rows[1].height = Mm(9)
        _widths(terms, [30] * 6)
        _grid(terms)
        _para(doc, space_after=10)

        self._items(doc, [(bill.no, ln) for ln in bill.lines], with_billno=False)
        self._totals(doc, [
            ("ยอดรวม", fmt_money(bill.subtotal), True),
            ("ภาษีขาย", "", False),
            ("การขนส่งและการจัดการ", date_th, False),
            ("วันครบกำหนดทั้งหมด", "", False),
        ])
        if billing.get("print_bahttext"):
            _para(doc, f"({bahttext(bill.subtotal)})", RIGHT)
        _para(doc, "ขอบคุณที่ร่วมทำธุรกิจกับเรา!", CENTER, bold=True, space_after=10).paragraph_format.space_before = Pt(16)
        signs = doc.add_table(rows=2, cols=3)
        _cell(signs.rows[0].cells[0], "ผู้จ่าย")
        _cell(signs.rows[0].cells[2], "ผู้รับ", RIGHT)
        _cell(signs.rows[1].cells[0], "______________________________")
        _cell(signs.rows[1].cells[2], "______________________________", RIGHT)
        _widths(signs, [60, 62, 60])

    def _header(self, doc, title: str, meta: list[str]) -> None:
        shop = self.settings["shop"]
        t = doc.add_table(rows=1, cols=2)
        left, right = t.rows[0].cells
        _cell(left, shop["name"], bold=True, size=13)
        for line in shop.get("address_lines", []):
            _para(left, str(line))
        _cell(right, title, RIGHT, bold=True, size=22)
        right.paragraphs[0].runs[0].font.color.rgb = GREY
        right.paragraphs[0].paragraph_format.space_after = Pt(18)
        for line in meta:
            _para(right, line, RIGHT)
        _widths(t, [90, 92])

    def _parties(self, doc) -> None:
        customer = self.settings["customer"]
        ship_to = customer.get("ship_to") or customer
        _para(doc, space_after=12)
        t = doc.add_table(rows=1, cols=2)
        for cell, label, who in zip(t.rows[0].cells, ["ถึง:", "จัดส่งไปยัง:"], [customer, ship_to]):
            _cell(cell, label, bold=True)
            _para(cell, who["name"])
            for line in who.get("address_lines", []):
                _para(cell, str(line))
            if who.get("tax_id"):
                _para(cell, f"เลขประจำตัวผู้เสียภาษี : {who['tax_id']}")
        _widths(t, [91, 91])

    def _items(self, doc, rows, with_billno: bool) -> None:
        head = (["ใบเสร็จ"] if with_billno else []) + ["จำนวน", "คำอธิบาย", "ราคาต่อหน่วย", "ผลรวม"]
        widths = ([16, 22, 100, 22, 22] if with_billno else [26, 112, 22, 22])
        t = doc.add_table(rows=1, cols=len(head))
        for c, h in zip(t.rows[0].cells, head):
            _cell(c, h, CENTER, bold=True)
        _repeat_header(t.rows[0])
        for billno, ln in rows:
            cells = iter(t.add_row().cells)
            if with_billno:
                _cell(next(cells), f"#{billno}", CENTER)
            _cell(next(cells), ln.qty_display, CENTER)
            _cell(next(cells), ln.description)
            _cell(next(cells), fmt_money(ln.unit_price), CENTER)
            _cell(next(cells), fmt_money(ln.amount), RIGHT)
        _widths(t, widths)
        _grid(t)

    def _totals(self, doc, rows: list[tuple[str, str, bool]]) -> None:
        t = doc.add_table(rows=0, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.RIGHT
        for label, value, strong in rows:
            label_cell, value_cell = t.add_row().cells
            _cell(label_cell, label, RIGHT, bold=strong)
            _cell(value_cell, value, RIGHT, bold=strong)
            _box(value_cell)
        _widths(t, [70, 28])

    def _save(self, doc, out: Path) -> Path:
        out.parent.mkdir(parents=True, exist_ok=True)
        doc.save(out)
        return out


def _box(cell) -> None:
    """เส้นรอบช่องตัวเลข (ช่องชื่อไม่มีเส้น เหมือน .totals ใน receipt.css)"""
    tcpr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "444444")
        borders.append(el)
    tcpr.append(borders)
