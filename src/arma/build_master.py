"""ขั้นที่ 0 — ตารางราคากลาง.docx → master/products.xlsx และ seed master/aliases.xlsx"""

from __future__ import annotations

import re
import zlib
from datetime import date
from pathlib import Path

import docx
import pandas as pd

from .normalize import Normalizer, norm_key

PRODUCT_COLUMNS = [
    "product_id", "category", "receipt_name", "color", "print_color", "unit_price",
    "model", "brand", "part", "logo", "size", "style", "thickness", "active", "valid_from",
]
ALIAS_COLUMNS = ["alias_key", "product_id", "example_text", "source", "confirmed_at"]

CATEGORY_BY_HEADING = {"เลเซอร์คัทอะคริลิค": "laser_acrylic", "3D Printing": "3d_print"}
CATEGORY_PREFIX = {"laser_acrylic": "LA", "3d_print": "3D", "service": "SV"}

COLOR_CODE = {
    "ดำเงา": "BLKGLOSS", "ดำด้าน": "BLKMATTE", "แดง": "RED", "ขาว": "WHITE", "เงิน": "SILVER",
    "ทอง": "GOLD", "เคฟล่าเงา": "KEVGLOSS", "เคฟล่าด้าน": "KEVMATTE", "ใสมีลาย": "CLEARPAT",
    "ใสไม่มีลาย": "CLEAR", "ดำเงา-แดง": "BLK-RED", "น้ำเงิน-แดง": "BLU-RED", "ดำเงา-เงิน": "BLK-SILVER",
}
PART_CODE = {
    "ฝากระโปรงหน้า": "HOOD", "ฝากระโปรงท้าย": "TRUNK", "ท้ายกระบะ": "REAR", "กระจังหน้า": "GRILLE",
    "มือจับท้าย": "HANDLE", "พรบ": "PRB", "ข้างรถ": "SIDE", "หัวเกียร์": "GEARKNOB",
    "กรองแอร์": "AIRFILTER", "ตะแกรงแตร": "HORNGRILLE", "ตะแกรง": "MESH", "แก้ว": "GLASS",
    "หัวชาร์จ": "CHARGER", "easy pass": "EASYPASS", "ตกแต่ง": "DECOR",
}
SIZE_CODE = {"ใหญ่": "L", "เล็ก": "S"}
# สีที่เป็นเนื้อสินค้า ไม่ต้องพิมพ์ "สีxxx" ต่อท้ายชื่อในใบเสร็จ
NO_PRINT_COLORS = {"ใสมีลาย", "ใสไม่มีลาย"}
# สินค้าที่ใบเสร็จจริง (#330–#336) ไม่พิมพ์สี แม้ลูกค้าระบุสีมา
NO_PRINT_COLOR_CATEGORIES = {"3d_print"}
NO_PRINT_COLOR_NAMES = ("มือจับท้าย",)


def prints_color(category: str, name: str, color: str) -> bool:
    return (
        bool(color)
        and color not in NO_PRINT_COLORS
        and "-" not in color
        and category not in NO_PRINT_COLOR_CATEGORIES
        and not any(k in name for k in NO_PRINT_COLOR_NAMES)
    )

# คำเรียกเริ่มต้นจากข้อมูลจริง พ.ค.–มิ.ย. 2569 (แผนข้อ 1.6): ข้อความ → (receipt_name, สี)
SEED_ALIASES = [
    ("กระจังหน้า fortuner gr", "FORTUNER LEGENDER โลโก้ GRเล็ก กระจังหน้า STYLE 2", ""),
    ("gr camry กระจัง", "TOYOTA CAMRY โลโก้ GRเล็ก กระจังหน้า STYLE 2", ""),
    ("กระจังหน้า gr travo", "TOYOTA TRAVO โลโก้ GRเล็ก กระจังหน้า STYLE 2", ""),
    ("gr กระจังหน้า yaris cross", "YARIS CROSS 2023 โลโก้ GRเล็ก กระจังหน้า STYLE 2", ""),
    ("โลโก้ กระจังหน้า yaris ativ gr", "YARIS ATIV โลโก้ GRเล็ก กระจังหน้า STYLE 2", ""),
]


def _cell(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def read_price_docx(path: Path) -> pd.DataFrame:
    """อ่านทุกตารางใน docx → 1 แถวต่อ (รายการ, กลุ่มสี) ตามต้นฉบับ"""
    doc = docx.Document(str(path))
    headings = [p.text for p in doc.paragraphs if p.text.strip().startswith("รายการ")]
    rows = []
    for i, table in enumerate(doc.tables):
        heading = headings[i] if i < len(headings) else ""
        category = next((c for k, c in CATEGORY_BY_HEADING.items() if k in heading), "laser_acrylic")
        header = [_cell(c.text) for c in table.rows[0].cells]
        col = {name: header.index(name) for name in ("รายการ", "สี", "ราคาต่อหน่วย")}
        for r in table.rows[1:]:
            cells = [_cell(c.text) for c in r.cells]
            name = cells[col["รายการ"]]
            if not name:
                continue
            price = re.sub(r"[^\d.]", "", cells[col["ราคาต่อหน่วย"]])
            rows.append({
                "category": category,
                "receipt_name": name,
                "color_group": cells[col["สี"]],
                "unit_price": float(price) if price else None,
            })
    return pd.DataFrame(rows)


def split_colors(group: str, category: str) -> list[str]:
    """'ดำเงา แดง ขาว' → 3 สี ; 'ดำเงา-แดง' = สีทูโทนสีเดียว ; '-' หรือว่าง = ไม่มีสี

    งาน 3D สีไม่มีผลต่อราคา จึงเพิ่มแถว "ไม่ระบุสี" ไว้ด้วย (ชื่อในใบเสร็จไม่มีสี)
    """
    group = group.strip()
    if group in ("", "-"):
        return [""]
    colors = group.split()
    return colors + [""] if category == "3d_print" else colors


def _product_id(category: str, name: str, attrs, color: str, has_colors: bool, taken: set[str]) -> str:
    ascii_part = re.sub(r"[^A-Za-z0-9 ]", " ", re.sub(r"[฀-๿]+", " ", name))
    tokens = [t.upper() for t in ascii_part.split()]
    tokens = [t for t in tokens if t not in {"MM"}]
    if re.search(r"\d\s*mm", name, re.I):
        tokens = [t for t in tokens if not t.isdigit() or not re.search(rf"{t}\s*mm", name, re.I)]
        tokens += [m.replace(" ", "").upper() for m in re.findall(r"\d\s*mm", name, re.I)]
    tokens += [PART_CODE[p] for p in sorted(attrs["part"]) if p in PART_CODE]
    tokens += [SIZE_CODE[s] for s in sorted(attrs["size"]) if s in SIZE_CODE]
    if color:
        tokens.append(COLOR_CODE.get(color, f"C{zlib.crc32(color.encode()) % 10000:04d}"))
    elif has_colors:
        tokens.append("ANY")
    base = "-".join([CATEGORY_PREFIX[category], *dict.fromkeys(tokens)])
    pid, n = base, 2
    while pid in taken:
        pid, n = f"{base}-{n}", n + 1
    taken.add(pid)
    return pid


def build_products(docx_path: Path, normalizer: Normalizer) -> pd.DataFrame:
    raw = read_price_docx(docx_path)
    taken: set[str] = set()
    rows = []
    for rec in raw.itertuples(index=False):
        attrs = normalizer.attributes(rec.receipt_name, apply_rules=False)
        colors = split_colors(rec.color_group, rec.category)
        for color in colors:
            rows.append({
                "product_id": _product_id(rec.category, rec.receipt_name, attrs, color, len(colors) > 1, taken),
                "category": rec.category,
                "receipt_name": rec.receipt_name,
                "color": color,
                "print_color": prints_color(rec.category, rec.receipt_name, color),
                "unit_price": rec.unit_price,
                **{a: "/".join(sorted(attrs[a])) for a in ("model", "brand", "part", "logo", "size", "style", "thickness")},
                "active": True,
                "valid_from": "",
            })
    # ค่าบริการ: ราคาแปรผัน อ่านจากข้อความ
    for pid, name, part in [
        ("SV-SHIPPING", "ค่าบริการจัดส่ง", "ค่าส่ง"),
        ("SV-DESIGN", "งานออกแบบและสั่งทำไฟล์", "ค่าออกแบบ"),
    ]:
        rows.append({
            "product_id": pid, "category": "service", "receipt_name": name,
            "color": "", "print_color": False, "unit_price": 0, "model": "", "brand": "",
            "part": part, "logo": "", "size": "", "style": "", "thickness": "",
            "active": True, "valid_from": "",
        })
    return pd.DataFrame(rows, columns=PRODUCT_COLUMNS)


def build_seed_aliases(products: pd.DataFrame, normalizer: Normalizer) -> pd.DataFrame:
    rows = []
    for text, name, color in SEED_ALIASES:
        hit = products[(products.receipt_name == name) & (products.color == color)]
        if hit.empty:
            raise ValueError(f"seed alias ชี้ไปสินค้าที่ไม่มี: {name} / {color}")
        rows.append({
            "alias_key": norm_key(normalizer.attributes(text)),
            "product_id": hit.iloc[0].product_id,
            "example_text": text,
            "source": "seed",
            "confirmed_at": date.today().isoformat(),
        })
    return pd.DataFrame(rows, columns=ALIAS_COLUMNS)


def write_xlsx(df: pd.DataFrame, path: Path, sheet: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        df.to_excel(xw, sheet_name=sheet, index=False)
        ws = xw.sheets[sheet]
        for col in ws.columns:
            width = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 70)
        ws.freeze_panes = "A2"
